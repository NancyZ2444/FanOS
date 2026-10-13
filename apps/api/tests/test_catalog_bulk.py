import csv
import io
from datetime import datetime

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.admin_models import AdminImportTask, ContentMetadata
from app.core.database import SessionLocal
from app.core.security import create_token
from app.models import Group, User
from app.modules.admin.catalog_service import MODELS

ROOT = "/api/v1/admin/console"
KINDS = list(MODELS)


def sheet(rows, fmt="csv"):
    fields = list(rows[0])
    if fmt == "csv":
        stream = io.StringIO()
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)
        return ("records.csv", stream.getvalue().encode("utf-8-sig"), "text/csv")
    workbook = Workbook()
    workbook.active.append(fields)
    for row in rows:
        workbook.active.append([row.get(field) for field in fields])
    stream = io.BytesIO()
    workbook.save(stream)
    return (
        "records.xlsx",
        stream.getvalue(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def preview(client, headers, kind, rows, fmt="csv"):
    result = client.post(
        ROOT + "/imports", headers=headers, data={"kind": kind}, files={"file": sheet(rows, fmt)}
    )
    assert result.status_code == 200, result.text
    return result.json()["data"]


def confirm(client, headers, task, skip=False):
    return client.post(
        ROOT + "/imports/" + task["id"] + "/confirm", headers=headers, json={"skip_invalid": skip}
    )


def create(client, headers, kind, name):
    result = client.post(
        ROOT + "/entities/" + kind,
        headers=headers,
        json={
            "name": name,
            "group_id": "group-lumina",
            **({"city": "Seoul"} if kind == "events" else {}),
        },
    )
    assert result.status_code == 200, result.text
    return result.json()["data"]["id"]


def batch(client, headers, kind, ids, action):
    return client.post(
        ROOT + "/entities/" + kind + "/batch", headers=headers, json={"ids": ids, "action": action}
    )


@pytest.mark.parametrize("kind", KINDS)
@pytest.mark.parametrize("fmt", ["csv", "xlsx"])
def test_import_each_entity_type_is_preview_only_then_persistent_and_idempotent(
    client, admin_headers, kind, fmt
):
    row = {
        "name": "Imported entity",
        "release_date": datetime(2026, 10, 9) if fmt == "xlsx" else "2026-10-09",
        "tags": "测试|新发行",
    }
    if kind not in {"groups", "others"}:
        row["group_name"] = "LUMINA"
    if kind == "events":
        row["city"] = "Seoul"
    if kind not in {"groups", "idols", "others"}:
        row["idol_name"] = "Aria"
    if kind == "others":
        row["start_date"] = row.pop("release_date")
    task = preview(client, admin_headers, kind, [row], fmt)
    assert task["kind"] == kind
    assert task["errors"] == []
    assert (
        task["rows"][0]["data"]["start_date" if kind == "others" else "release_date"]
        == "2026-10-09"
    )
    with SessionLocal() as db:
        assert db.scalar(select(MODELS[kind]).where(MODELS[kind].name == "Imported entity")) is None
    response = confirm(client, admin_headers, task)
    assert response.status_code == 200, response.text
    created_id = response.json()["data"]["rows"][0]["created_id"]
    assert (
        confirm(client, admin_headers, task).json()["data"]["rows"][0]["created_id"] == created_id
    )
    with SessionLocal() as db:
        assert db.get(MODELS[kind], created_id).name == "Imported entity"
        meta = db.scalar(select(ContentMetadata).where(ContentMetadata.entity_id == created_id))
        assert meta.status == "DRAFT"
        assert meta.data["tags"] == ["测试", "新发行"]
    history = client.get(ROOT + "/imports?kind=" + kind, headers=admin_headers).json()["data"][
        "items"
    ]
    assert len(history) == 1 and history[0]["kind"] == kind
    duplicate = preview(client, admin_headers, kind, [row], fmt)
    assert duplicate["rows"] == []
    assert duplicate["errors"][0]["field"] == "name"


@pytest.mark.parametrize("kind", KINDS)
def test_entity_batch_lifecycle_and_atomic_rollback(client, admin_headers, kind):
    ids = [
        create(client, admin_headers, kind, "Bulk A"),
        create(client, admin_headers, kind, "Bulk B"),
    ]
    for action, status in [
        ("submit", "PENDING"),
        ("publish", "PUBLISHED"),
        ("offline", "OFFLINE"),
        ("archive", "ARCHIVED"),
        ("restore", "DRAFT"),
    ]:
        response = batch(client, admin_headers, kind, ids, action)
        assert response.status_code == 200, response.text
        assert {r["status"] for r in response.json()["data"]["items"]} == {status}
    assert batch(client, admin_headers, kind, [ids[0]], "submit").status_code == 200
    assert batch(client, admin_headers, kind, [ids[0]], "publish").status_code == 200
    failure = batch(client, admin_headers, kind, ids, "offline")
    assert failure.status_code == 409
    assert ids[1] in failure.json()["message"]
    with SessionLocal() as db:
        first = db.scalar(select(ContentMetadata).where(ContentMetadata.entity_id == ids[0]))
        assert first.status == "PUBLISHED"
    missing = batch(client, admin_headers, kind, [ids[0], "not-found"], "offline")
    assert missing.status_code == 404
    with SessionLocal() as db:
        assert (
            db.scalar(select(ContentMetadata).where(ContentMetadata.entity_id == ids[0])).status
            == "PUBLISHED"
        )


def test_import_reports_duplicates_missing_links_and_formats_per_row(client, admin_headers):
    rows = [
        {"name": "New album", "group_name": "LUMINA", "idol_id": "", "release_date": "2026-10-09"},
        {"name": "New album", "group_name": "LUMINA", "idol_id": "", "release_date": "2026-10-09"},
        {"name": "Bad group", "group_name": "Unknown", "idol_id": "", "release_date": ""},
        {"name": "Bad idol", "group_name": "LUMINA", "idol_id": "missing-idol", "release_date": ""},
        {"name": "Bad date", "group_name": "LUMINA", "idol_id": "", "release_date": "tomorrow"},
    ]
    task = preview(client, admin_headers, "albums", rows)
    assert len(task["rows"]) == 1
    assert {(e["row"], e["field"]) for e in task["errors"]} == {
        (3, "name"),
        (4, "group_name"),
        (5, "idol_id"),
        (6, "release_date"),
    }
    assert confirm(client, admin_headers, task).status_code == 422
    assert confirm(client, admin_headers, task, True).json()["data"]["imported_count"] == 1


def test_confirmation_revalidates_and_rolls_back_if_database_changed(client, admin_headers):
    task = preview(
        client, admin_headers, "groups", [{"name": "First import"}, {"name": "Second import"}]
    )
    create(client, admin_headers, "groups", "Second import")
    failed = confirm(client, admin_headers, task)
    assert failed.status_code == 422
    assert failed.json()["data"]["details"][0]["row"] == 3
    with SessionLocal() as db:
        assert db.scalar(select(Group).where(Group.name == "First import")) is None
        assert db.get(AdminImportTask, task["id"]).status == "VALIDATED"


def test_import_and_batch_permissions_and_limits(client, admin_headers, auth_headers):
    with SessionLocal() as db:
        db.add_all(
            [
                User(id="bulk-content", display_name="Content", role="content_operator"),
                User(id="bulk-ops", display_name="Ops", role="operation_admin"),
            ]
        )
        db.commit()
    content = {"Authorization": "Bearer " + create_token("bulk-content", "access")}
    ops = {"Authorization": "Bearer " + create_token("bulk-ops", "access")}
    ids = [create(client, content, "groups", "Allowed group")]
    assert batch(client, content, "groups", ids, "submit").status_code == 200
    assert batch(client, content, "groups", ids, "publish").status_code == 403
    assert batch(client, ops, "groups", ids, "archive").status_code == 403
    assert batch(client, auth_headers, "groups", ids, "archive").status_code == 403
    assert batch(client, admin_headers, "groups", [], "archive").status_code == 422
    assert batch(client, admin_headers, "groups", ["x"] * 101, "archive").status_code == 422
    for endpoint in ["/imports", "/imports/template?kind=groups"]:
        assert client.get(ROOT + endpoint, headers=ops).status_code == 403
    assert (
        client.post(
            ROOT + "/imports",
            headers=ops,
            data={"kind": "groups"},
            files={"file": sheet([{"name": "No permission"}])},
        ).status_code
        == 403
    )
    assert (
        client.post(
            ROOT + "/imports",
            headers=admin_headers,
            data={"kind": "unsupported"},
            files={"file": sheet([{"name": "Wrong kind"}])},
        ).status_code
        == 422
    )
    task = preview(client, content, "groups", [{"name": "Content import"}])
    assert confirm(client, ops, task).status_code == 403
    assert confirm(client, content, task).status_code == 200


@pytest.mark.parametrize("kind", ["cards", *KINDS])
def test_templates_csv_and_excel_are_specific_to_kind(client, admin_headers, kind):
    csv_res = client.get(ROOT + f"/imports/template?kind={kind}&format=csv", headers=admin_headers)
    xlsx_res = client.get(
        ROOT + f"/imports/template?kind={kind}&format=xlsx", headers=admin_headers
    )
    assert csv_res.status_code == xlsx_res.status_code == 200
    fields = next(csv.reader(io.StringIO(csv_res.content.decode("utf-8-sig"))))
    workbook = load_workbook(io.BytesIO(xlsx_res.content))
    assert list(next(workbook.active.values)) == fields
    assert fields[0] == ("group_id" if kind == "cards" else "name")
    if kind not in {"cards", "groups", "others"}:
        assert "group_name" in fields
    workbook.close()
