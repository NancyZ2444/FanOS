from __future__ import annotations

import csv
import io
from typing import Literal

from fastapi import APIRouter, Depends, File, Form, UploadFile
from fastapi.responses import Response
from PIL import Image, ImageOps, UnidentifiedImageError
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.admin_models import AdminImportTask, CardImageVersion
from app.common.response import DomainError, success
from app.core.database import get_db
from app.core.security import current_admin
from app.integrations.storage import get_storage
from app.models import Card, OperationLog, User
from app.modules.admin import catalog_service as service
from app.modules.admin import import_service
from app.modules.admin.console_schemas import (
    ActionRequest,
    BatchAction,
    CardDraft,
    EntityDraft,
    ImportConfirm,
    ImportKind,
)

router = APIRouter(prefix="/admin/console", tags=["Admin Console"])


@router.get("/me")
def me(actor: User = Depends(current_admin)) -> dict:
    return success({"id": actor.id, "name": actor.display_name, "role": actor.role})


@router.get("/catalog")
def catalog(actor: User = Depends(current_admin), db: Session = Depends(get_db)) -> dict:
    # All staff can read catalog choices for notification links / card pools.
    result = service.all_entities(db)
    result["cards"] = [
        service.admin_card(db, c)
        for c in db.scalars(select(Card).order_by(Card.updated_at.desc())).all()
    ]
    return success(result)


@router.post("/cards")
def create_card(
    payload: CardDraft, actor: User = Depends(current_admin), db: Session = Depends(get_db)
) -> dict:
    row = service.save_card(db, actor, payload)
    db.commit()
    return success(service.admin_card(db, row))


@router.post("/cards/crop-preview")
async def card_crop_preview(
    file: UploadFile = File(...), actor: User = Depends(current_admin)
) -> dict:
    from starlette.concurrency import run_in_threadpool

    from .card_cropping import crop_preview

    service.permit(actor, "catalog", "edit")
    raw = await file.read(10 * 1024 * 1024 + 1)
    if len(raw) > 10 * 1024 * 1024:
        raise DomainError("FILE_TOO_LARGE", "图片不能超过 10 MB", 413)
    return success(await run_in_threadpool(crop_preview, raw))


@router.post("/cards/recognize")
async def recognize_card(
    file: UploadFile = File(...),
    crop: str | None = Form(
        None, description="已确认的裁剪四角 JSON：{points:[{x,y},…]}，坐标 0–1，顺时针"
    ),
    actor: User = Depends(current_admin),
    db: Session = Depends(get_db),
) -> dict:
    from starlette.concurrency import run_in_threadpool

    from .card_recognition import recognize

    service.permit(actor, "catalog", "edit")
    raw = await file.read(10 * 1024 * 1024 + 1)
    if len(raw) > 10 * 1024 * 1024:
        raise DomainError("FILE_TOO_LARGE", "图片不能超过 10 MB", 413)
    return success(await run_in_threadpool(recognize, db, actor, raw, crop))


@router.put("/cards/{card_id}")
def update_card(
    card_id: str,
    payload: CardDraft,
    actor: User = Depends(current_admin),
    db: Session = Depends(get_db),
) -> dict:
    row = service.save_card(db, actor, payload, card_id)
    db.commit()
    return success(service.admin_card(db, row))


@router.post("/cards/batch")
def batch_cards(
    payload: BatchAction, actor: User = Depends(current_admin), db: Session = Depends(get_db)
) -> dict:
    # One transaction: a failed row never leaves half of the selection published.
    service.permit(actor, "catalog", payload.action)
    try:
        results = [
            service.delete_card(db, actor, card_id)
            if payload.action == "delete"
            else service.transition_card(db, actor, card_id, payload.action)
            for card_id in dict.fromkeys(payload.ids)
        ]
    except DomainError as exc:
        raise DomainError(
            exc.error_code, f"批量操作失败：{exc.message}；整批未执行", exc.status_code
        ) from exc
    db.commit()
    return success({"items": results})


@router.delete("/cards/{card_id}")
def delete_card(
    card_id: str,
    actor: User = Depends(current_admin),
    db: Session = Depends(get_db),
) -> dict:
    result = service.delete_card(db, actor, card_id)
    db.commit()
    return success(result)


@router.post("/cards/{card_id}/actions")
def card_action(
    card_id: str,
    payload: ActionRequest,
    actor: User = Depends(current_admin),
    db: Session = Depends(get_db),
) -> dict:
    result = service.transition_card(db, actor, card_id, payload.action, payload.reason)
    db.commit()
    return success(result)


@router.get("/cards/{card_id}/source-history")
def source_history(
    card_id: str, actor: User = Depends(current_admin), db: Session = Depends(get_db)
) -> dict:
    service.permit(actor, "sources")
    return success(
        {
            "items": [
                service.row_data(row)
                for row in db.scalars(
                    select(CardImageVersion)
                    .where(CardImageVersion.card_id == card_id)
                    .order_by(CardImageVersion.created_at.desc())
                ).all()
            ]
        }
    )


@router.post("/entities/{kind}")
def create_entity(
    kind: str,
    payload: EntityDraft,
    actor: User = Depends(current_admin),
    db: Session = Depends(get_db),
) -> dict:
    result = service.save_entity(db, actor, kind, payload)
    db.commit()
    return success(result)


@router.put("/entities/{kind}/{item_id}")
def update_entity(
    kind: str,
    item_id: str,
    payload: EntityDraft,
    actor: User = Depends(current_admin),
    db: Session = Depends(get_db),
) -> dict:
    result = service.save_entity(db, actor, kind, payload, item_id)
    db.commit()
    return success(result)


@router.post("/entities/{kind}/batch")
def batch_entities(
    kind: str,
    payload: BatchAction,
    actor: User = Depends(current_admin),
    db: Session = Depends(get_db),
) -> dict:
    service.permit(actor, "catalog", payload.action)
    if kind not in service.MODELS:
        raise DomainError("INVALID_ENTITY_TYPE", "不支持的图鉴类型", 404)
    results = []
    for item_id in dict.fromkeys(payload.ids):
        try:
            results.append(
                service.delete_entity(db, actor, kind, item_id)
                if payload.action == "delete"
                else service.transition_entity(db, actor, kind, item_id, payload.action)
            )
        except DomainError as exc:
            # The dependency closes this uncommitted session: all earlier rows roll back.
            raise DomainError(
                exc.error_code, f"记录 {item_id}：{exc.message}；整批未执行", exc.status_code
            ) from exc
    db.commit()
    return success({"items": results})


@router.delete("/entities/{kind}/{item_id}")
def delete_entity(
    kind: str,
    item_id: str,
    actor: User = Depends(current_admin),
    db: Session = Depends(get_db),
) -> dict:
    result = service.delete_entity(db, actor, kind, item_id)
    db.commit()
    return success(result)


@router.post("/entities/{kind}/{item_id}/actions")
def entity_action(
    kind: str,
    item_id: str,
    payload: ActionRequest,
    actor: User = Depends(current_admin),
    db: Session = Depends(get_db),
) -> dict:
    result = service.transition_entity(db, actor, kind, item_id, payload.action, payload.reason)
    db.commit()
    return success(result)


@router.post("/uploads")
async def upload(
    file: UploadFile = File(...),
    crop: str | None = Form(None, description="可选小卡裁剪四角 JSON，坐标 0–1，顺时针"),
    actor: User = Depends(current_admin),
) -> dict:
    from .card_cropping import apply_crop, parse_crop

    confirmed = None
    if crop is not None:
        service.permit(actor, "catalog", "edit")
        confirmed = parse_crop(crop)
    raw = await file.read(10 * 1024 * 1024 + 1)
    if len(raw) > 10 * 1024 * 1024:
        raise DomainError("FILE_TOO_LARGE", "图片不能超过 10 MB", 413)
    try:
        image = Image.open(io.BytesIO(raw))
        if image.format not in {"JPEG", "PNG", "WEBP"} or image.width * image.height > 40_000_000:
            raise ValueError("invalid image")
        image.load()
        image = ImageOps.exif_transpose(image).convert("RGB")
    except (UnidentifiedImageError, OSError, ValueError, Image.DecompressionBombError) as exc:
        raise DomainError("INVALID_IMAGE", "请上传有效的 JPG、PNG 或 WebP 图片", 422) from exc
    if confirmed is not None:
        image = apply_crop(image, confirmed)
    storage = get_storage()
    variants = {}
    for name, size in [("original", 5000), ("medium", 1200), ("thumbnail", 360)]:
        copy = image.copy()
        copy.thumbnail((size, size))
        buffer = io.BytesIO()
        copy.save(buffer, format="WEBP", quality=90)
        variants[name] = storage.save_bytes(buffer.getvalue(), "image.webp", f"catalog/{actor.id}")
    return success(
        {
            "url": variants["original"],
            "variants": variants,
            "width": image.width,
            "height": image.height,
        }
    )


@router.get("/imports")
def imports(
    kind: ImportKind | None = None,
    actor: User = Depends(current_admin),
    db: Session = Depends(get_db),
) -> dict:
    service.permit(actor, "catalog")
    return success(
        {
            "items": [
                service.row_data(row)
                for row in db.scalars(
                    select(AdminImportTask)
                    .where(AdminImportTask.kind == kind if kind else True)
                    .order_by(AdminImportTask.created_at.desc())
                ).all()
            ]
        }
    )


@router.get("/imports/template")
def import_template(
    kind: ImportKind = "cards",
    format: Literal["csv", "xlsx"] = "csv",
    actor: User = Depends(current_admin),
) -> Response:
    service.permit(actor, "catalog")
    fields = import_service.columns(kind)
    if format == "xlsx":
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = kind
        sheet.append(fields)
        sheet.freeze_panes = "A2"
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter

        for i, cell in enumerate(sheet[1], start=1):
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="8060C8")
            sheet.column_dimensions[get_column_letter(i)].width = max(20, len(cell.value) + 3)
        buffer = io.BytesIO()
        workbook.save(buffer)
        content = buffer.getvalue()
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        buffer = io.StringIO()
        csv.writer(buffer).writerow(fields)
        content = buffer.getvalue().encode("utf-8-sig")
        media_type = "text/csv; charset=utf-8"
    return Response(
        content,
        media_type=media_type,
        headers={
            "Content-Disposition": f'attachment; filename="FanOS-{kind}-template.{format}"',
            "Cache-Control": "no-store, max-age=0",
        },
    )


@router.post("/imports")
async def validate_import(
    file: UploadFile = File(...),
    kind: ImportKind = Form("cards"),
    actor: User = Depends(current_admin),
    db: Session = Depends(get_db),
) -> dict:
    service.permit(actor, "catalog", "edit")
    raw = await file.read(5 * 1024 * 1024 + 1)
    records = import_service.read_sheet(file.filename or "", raw, kind)
    valid_rows, errors = import_service.validate_rows(db, kind, records)
    task = AdminImportTask(
        filename=(file.filename or "import")[:240],
        kind=kind,
        rows=valid_rows,
        errors=errors,
        status="HAS_ERRORS" if errors else "VALIDATED",
        operator_id=actor.id,
    )
    db.add(task)
    db.flush()
    from fastapi.encoders import jsonable_encoder

    from app.modules.admin.catalog_extensions import CatalogRecord

    db.add(
        CatalogRecord(
            kind="import_input",
            entity_type="import",
            entity_id=task.id,
            data={"rows": jsonable_encoder(records), "skipped_rows": []},
        )
    )
    db.commit()
    db.refresh(task)
    return success(service.row_data(task))


@router.post("/imports/{task_id}/confirm")
def confirm_import(
    task_id: str,
    payload: ImportConfirm,
    actor: User = Depends(current_admin),
    db: Session = Depends(get_db),
) -> dict:
    service.permit(actor, "catalog", "edit")
    task = db.scalar(select(AdminImportTask).where(AdminImportTask.id == task_id).with_for_update())
    if not task:
        raise DomainError("IMPORT_NOT_FOUND", "导入任务不存在", 404)
    if task.status == "COMPLETED":
        return success(service.row_data(task))
    if task.errors and not payload.skip_invalid:
        raise DomainError("IMPORT_HAS_ERRORS", "请修复错误行，或明确选择跳过错误行", 422)
    if not task.rows:
        raise DomainError("NO_VALID_ROWS", "没有可导入的数据行", 422)
    combined_values = iter(())
    combined_groups = {}
    if task.kind == "groups_idols":
        sources = [
            {
                key: value
                for key, value in row["data"].items()
                if key in import_service.COMBINED_COLUMNS
            }
            for row in task.rows
        ]
        valid, errors = import_service.validate_combined(db, sources)
        service.fail_fields(
            [
                {
                    **error,
                    "row": task.rows[error["row"] - 2]["row"],
                    "message": (f"第 {task.rows[error['row'] - 2]['row']} 行：{error['message']}"),
                }
                for error in errors
            ]
        )
        combined_values = iter(row["data"] for row in valid)
    imported_rows = []
    for row in task.rows:
        if task.kind == "companies":
            from .company_import import create_company

            created_id = create_company(db, actor, row)
            imported_rows.append({**row, "created_id": created_id})
            continue
        if task.kind == "groups_idols":
            result = import_service.import_combined(
                db, actor, next(combined_values), combined_groups
            )
            imported_rows.append({**row, **result})
            continue
        if task.kind == "cards":
            draft = CardDraft.model_validate(row["data"])
            errors = service.validate_relations(db, draft, True)
        else:
            draft = EntityDraft.model_validate(row["data"])
            errors = import_service.relation_errors(db, task.kind, draft)
        service.fail_fields(
            [
                {**error, "row": row["row"], "message": f"第 {row['row']} 行：{error['message']}"}
                for error in errors
            ]
        )
        if task.kind == "cards":
            created_id = service.save_card(db, actor, draft).id
        elif task.kind == "albums":
            created_id, version_ids, mode = import_service.import_album(db, actor, draft)
            row = {**row, "created_version_ids": version_ids, "import_mode": mode}
        else:
            created_id = service.save_entity(db, actor, task.kind, draft)["id"]
        imported_rows.append({**row, "created_id": created_id})
    task.rows = imported_rows
    task.status = "COMPLETED"
    task.imported_count = len(task.rows)
    service.audit(
        db,
        actor,
        "catalog.import.confirm",
        "import",
        task.id,
        None,
        {
            "kind": task.kind,
            "imported_count": task.imported_count,
            "skipped": len({e["row"] for e in task.errors}),
        },
    )
    db.commit()
    db.refresh(task)
    return success(service.row_data(task))


@router.get("/logs")
def logs(actor: User = Depends(current_admin), db: Session = Depends(get_db)) -> dict:
    if actor.role != "admin":
        raise DomainError("PERMISSION_DENIED", "仅管理员可查看操作日志", 403)
    return success(
        {
            "items": [
                service.row_data(row)
                for row in db.scalars(
                    select(OperationLog).order_by(OperationLog.created_at.desc()).limit(500)
                ).all()
            ]
        }
    )
