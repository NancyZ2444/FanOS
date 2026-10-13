"""Catalog spreadsheet parsing and preflight; never writes catalog rows during validation."""

from __future__ import annotations

import csv
import io
from datetime import date, datetime
from itertools import islice
from pathlib import Path
from urllib.parse import urlsplit

from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.common.response import DomainError
from app.models import Group, GroupMembership, Idol
from app.modules.admin import catalog_service as service
from app.modules.admin.console_schemas import CardDraft, EntityDraft

from .catalog_companies import (
    company_name_key,
    ensure_company_relation,
    get_or_create_company,
    normalize_company_name,
    resolve_company_name,
)

CARD_COLUMNS = [
    "group_id",
    "idol_id",
    "source_entity_type",
    "source_entity_id",
    "card_type",
    "version",
    "edition",
    "rarity",
    "release_date",
    "image_url",
    "front_source_type",
    "front_source_url",
    "front_copyright_status",
    "back_image_url",
    "back_source_type",
    "back_source_url",
    "back_copyright_status",
    "is_group_pc",
    "is_double_sided",
    "is_official",
    "is_limited",
    "is_random",
    "has_ar",
    "ar_type",
    "ar_resource",
    "tags",
]
ENTITY_COLUMNS = [
    "name",
    "name_en",
    "name_ko",
    "name_translated",
    "description",
    "release_date",
    "image_url",
    "tags",
    "source_type",
    "source_url",
    "copyright_status",
    "copyright_note",
]
COMBINED_GROUP_FIELDS = {
    "group_code": "group_code",
    "group_name": "name",
    "group_name_en": "name_en",
    "group_name_ko": "name_ko",
    "group_name_translated": "name_translated",
    "group_company": "company",
    "group_description": "description",
    "group_debut_date": "release_date",
    "group_end_date": "end_date",
    "group_region": "region",
    "group_status": "business_status",
    "group_fandom_name": "fandom_name",
    "group_official_website": "official_website",
    "group_aliases": "aliases",
    "group_image_url": "image_url",
    "group_tags": "tags",
    "group_source_type": "source_type",
    "group_source_url": "source_url",
    "group_copyright_status": "copyright_status",
    "group_copyright_note": "copyright_note",
}
COMBINED_IDOL_FIELDS = {
    "stage_name": "name",
    "idol_name": "real_name",
    "idol_name_en": "name_en",
    "idol_name_ko": "name_ko",
    "idol_name_translated": "name_translated",
    "idol_company": "company",
    "idol_description": "description",
    "birthday": "release_date",
    "idol_region": "region",
    "idol_gender": "gender",
    "idol_aliases": "aliases",
    "idol_image_url": "image_url",
    "idol_tags": "tags",
    "idol_source_type": "source_type",
    "idol_source_url": "source_url",
    "idol_copyright_status": "copyright_status",
    "idol_copyright_note": "copyright_note",
}
COMBINED_COLUMNS = [
    "group_id",
    *COMBINED_GROUP_FIELDS,
    "idol_id",
    "idol_name",
    "stage_name",
    *[field for field in COMBINED_IDOL_FIELDS if field not in {"idol_name", "stage_name"}],
]
COMBINED_GROUP_PROFILE_FIELDS = ["group_id", *COMBINED_GROUP_FIELDS]


def columns(kind: str) -> list[str]:
    if kind == "others":
        return ["name", "original_name", "start_date", "end_date", "description", "tags"]
    if kind == "companies":
        from .company_import import COLUMNS

        return COLUMNS
    if kind == "groups_idols":
        return COMBINED_COLUMNS
    if kind == "cards":
        return CARD_COLUMNS
    if kind == "groups":
        return ["name", "company", *ENTITY_COLUMNS[1:]]
    extras = ["group_id", "group_name"]
    if kind != "idols":
        extras += ["idol_id", "idol_name", "type", "region"]
    if kind == "albums":
        extras += ["artist_names", "versions"]
    if kind == "events":
        extras += [
            "is_online",
            "city",
            "platform",
            "venue",
            "end_date",
            "parent_event_id",
            "album_id",
        ]
    return ["name", *extras, *ENTITY_COLUMNS[1:]]


def read_sheet(filename: str, raw: bytes, kind: str) -> list[dict]:
    if len(raw) > 5 * 1024 * 1024:
        raise DomainError("FILE_TOO_LARGE", "导入文件不能超过 5 MB", 413)
    try:
        suffix = Path(filename).suffix.lower()
        if suffix == ".csv":
            reader = csv.DictReader(io.StringIO(raw.decode("utf-8-sig")))
            fields = reader.fieldnames or []
            records = list(islice(reader, 2001))
        elif suffix == ".xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            try:
                rows = workbook.active.iter_rows(values_only=True)
                fields = [str(v or "") for v in next(rows)]
                records = [dict(zip(fields, row, strict=False)) for row in islice(rows, 2001)]
            finally:
                workbook.close()
        else:
            raise ValueError("unsupported file")
    except Exception as exc:
        raise DomainError(
            "INVALID_IMPORT_FILE", "无法读取文件，请使用 UTF-8 CSV 或 XLSX 模板", 422
        ) from exc
    if not fields or len(fields) != len(set(fields)) or any(not f for f in fields):
        raise DomainError("INVALID_HEADERS", "列名不能为空或重复，请下载对应类型模板", 422)
    if kind == "groups_idols":
        required = {"group_code", "group_name", "idol_name"}
        if not required.issubset(fields) or set(fields) - set(columns(kind)):
            raise DomainError("INVALID_HEADERS", "请使用组合 + 艺人联合导入模板", 422)
        if not records or len(records) > 2000:
            raise DomainError("INVALID_ROW_COUNT", "每次请导入 1–2000 行", 422)
        return records
    required = {"group_id", "source_entity_id", "image_url"} if kind == "cards" else {"name"}
    if kind not in {"cards", "groups", "idols", "companies", "others"} and not {
        "group_id",
        "group_name",
        "idol_id",
        "idol_name",
        "artists",
        "artist_names",
        "idol_names",
        "idol_ids",
    } & set(fields):
        raise DomainError("INVALID_HEADERS", "请提供组合或个人艺人归属", 422)
    if not required.issubset(fields):
        raise DomainError(
            "INVALID_HEADERS", "缺少列：" + "、".join(sorted(required - set(fields))), 422
        )
    allowed = set(columns(kind)) | set((CardDraft if kind == "cards" else EntityDraft).model_fields)
    if kind in {"companies", "others"}:
        allowed = set(columns(kind))
    if kind == "albums":
        allowed |= {"version", "idol_names", "idol_ids"}
    if set(fields) - allowed:
        raise DomainError(
            "INVALID_HEADERS", "当前类型不支持列：" + "、".join(sorted(set(fields) - allowed)), 422
        )
    if not records or len(records) > 2000:
        raise DomainError("INVALID_ROW_COUNT", "每次请导入 1–2000 行", 422)
    return records


def fail(field: str, message: str):
    service.fail_fields([{"field": field, "message": message}])


def artist_key(payload: EntityDraft):
    artists = payload.artists or (
        [{"type": "groups", "id": payload.group_id}] if payload.group_id else []
    )
    result = {(a.get("type"), a.get("id")) for a in artists}
    if payload.idol_id:
        result.add(("idols", payload.idol_id))
    return tuple(sorted(result))


def relation_errors(db: Session, kind: str, payload: EntityDraft) -> list[dict]:
    errors = []
    from app.modules.catalog.contributions import validate_other_fields
    from app.modules.catalog.tags import TAG_TYPES, validate_names

    try:
        if kind in TAG_TYPES:
            validate_names(db, TAG_TYPES[kind], payload.tags)
        if kind == "others":
            validate_other_fields(payload.model_dump())
    except DomainError as exc:
        errors.append(
            {
                "field": "tags" if exc.error_code == "TAG_DISABLED" else "data",
                "message": exc.message,
            }
        )
    if kind != "groups":
        if payload.group_id and not db.get(Group, payload.group_id):
            errors.append({"field": "group_id", "message": "组合不存在，请先导入组合或修正 ID"})
        if payload.idol_id and not db.get(Idol, payload.idol_id):
            errors.append({"field": "idol_id", "message": "艺人不存在，请修正 ID"})
        if kind not in {"idols", "others"} and not (
            payload.group_id or payload.idol_id or payload.artists
        ):
            errors.append({"field": "artists", "message": "请选择组合或个人归属"})
        for artist in payload.artists:
            artist_kind = artist.get("type")
            if artist_kind not in {"groups", "idols"} or not db.get(
                Group if artist_kind == "groups" else Idol, artist.get("id", "")
            ):
                errors.append({"field": "artists", "message": "参与组合或艺人不存在"})
    model = service.MODELS[kind]
    query = select(model).where(model.name == payload.name)
    if kind == "idols":
        if payload.group_id:
            query = query.join(GroupMembership, GroupMembership.idol_id == Idol.id).where(
                GroupMembership.group_id == payload.group_id
            )
        else:
            query = query.where(
                Idol.id.not_in(
                    select(GroupMembership.idol_id).where(
                        GroupMembership.status.in_(["active", "inactive"])
                    )
                )
            )
    existing = db.scalars(query).all()
    if kind not in {"groups", "idols", "others"}:
        from .catalog_extensions import actor_names

        existing = [
            r
            for r in existing
            if artist_key(EntityDraft(name=r.name, artists=actor_names(db, kind, r)))
            == artist_key(payload)
        ]
    if existing and kind == "albums" and payload.versions:
        from .album_versions import album_names, version_key

        if len(existing) > 1:
            errors.append(
                {"field": "name", "message": "匹配到多个同名同艺人专辑，请先修正重复父专辑"}
            )
        else:
            known = {version_key(n) for n in album_names(db, existing[0])}
            repeated = [n for n in payload.versions if version_key(n) in known]
            if repeated:
                errors.append(
                    {"field": "versions", "message": "版本已存在：" + "、".join(repeated)}
                )
    elif existing:
        errors.append(
            {"field": "name", "message": "已存在同组合同名记录（含草稿或归档），请在列表编辑原记录"}
        )
    return errors


def matching_album(db: Session, payload: EntityDraft):
    from app.models import CatalogAlbum

    from .catalog_extensions import actor_names

    matches = [
        a
        for a in db.scalars(select(CatalogAlbum).where(CatalogAlbum.name == payload.name))
        if artist_key(EntityDraft(name=a.name, artists=actor_names(db, "albums", a)))
        == artist_key(payload)
    ]
    if len(matches) > 1:
        fail("name", "匹配到多个同名同艺人专辑，请先修正重复父专辑")
    return matches[0] if matches else None


def import_album(db, actor, payload):
    """Version rows append to a single parent, without replacing parent attributes."""
    from .album_versions import album_records, ensure_versions
    from .catalog_service import audit, entity_payload

    album = matching_album(db, payload)
    if album is None:
        result = service.save_entity(db, actor, "albums", payload)
        version_ids = [
            r.id for r in album_records(db, result["id"]) if r.data["name"] in payload.versions
        ]
        return result["id"], version_ids, "created"
    before = entity_payload(db, "albums", album)
    ids, _ = ensure_versions(db, actor, album, payload.versions)
    db.flush()
    audit(
        db,
        actor,
        "catalog.album.append_versions",
        "albums",
        album.id,
        before,
        entity_payload(db, "albums", album),
    )
    return album.id, ids, "appended_versions"


def resolve_artist(db, field, name, group_id="", only_idols=False):
    matches = []
    if not only_idols:
        matches.extend(
            {"type": "groups", "id": g.id}
            for g in db.scalars(select(Group).where(Group.name == name))
        )
    idols = list({i.id: i for i in db.scalars(select(Idol).where(Idol.name == name))}.values())
    if len(idols) > 1 and group_id:
        linked = db.scalars(
            select(Idol)
            .join(GroupMembership)
            .where(Idol.name == name, GroupMembership.group_id == group_id)
        ).all()
        if linked:
            idols = list({i.id: i for i in linked}.values())
    matches.extend({"type": "idols", "id": i.id} for i in idols)
    if len(matches) != 1:
        fail(field, f"艺人或组合名称「{name}」不存在或不唯一，请使用 artists 中的准确 ID")
    return matches[0]


def entity_draft(db: Session, kind: str, value: dict) -> EntityDraft:
    if kind == "others":
        from app.modules.catalog.contributions import validate_other_fields

        validate_other_fields(value)
        return EntityDraft.model_validate(value)
    group_name = value.pop("group_name", "")
    if kind != "groups":
        group_id = value.get("group_id", "")
        group = db.get(Group, group_id) if group_id else None
        if group_id and not group:
            fail("group_id", "组合 ID 不存在，请先导入组合")
        if group_name:
            if group and group.name != group_name:
                fail("group_name", "组合名称与 group_id 不一致")
            if not group:
                matches = db.scalars(select(Group).where(Group.name == group_name)).all()
                if len(matches) != 1:
                    fail("group_name", "组合名称不存在或不唯一，请填写准确的 group_id")
                group = matches[0]
        value["group_id"] = group.id if group else ""
    idol_name = value.pop("idol_name", "")
    artist_names = value.pop("artist_names", "")
    idol_names = value.pop("idol_names", "")
    idol_ids = value.pop("idol_ids", "")
    if kind == "albums" and (artist_names or idol_names or idol_ids or "|" in idol_name):
        if value.get("artists"):
            EntityDraft.model_validate(
                {"name": value.get("name") or "validation", "artists": value["artists"]}
            )
            if any(not a.get("type") or not a.get("id") for a in value["artists"]):
                fail("artists", "参与艺人须提供 type 与 id")
        artists = list(value.get("artists") or [])
        if value.get("group_id"):
            artists.append({"type": "groups", "id": value["group_id"]})
        for field, text, only in [
            ("artist_names", artist_names, False),
            ("idol_names", idol_names, True),
        ]:
            for name in [n.strip() for n in text.split("|") if n.strip()]:
                artists.append(resolve_artist(db, field, name, value.get("group_id", ""), only))
        if idol_name:
            for name in [n.strip() for n in idol_name.split("|") if n.strip()]:
                artists.append(
                    resolve_artist(db, "idol_name", name, value.get("group_id", ""), True)
                )
            idol_name = ""
        for iid in [i.strip() for i in idol_ids.split("|") if i.strip()]:
            if not db.get(Idol, iid):
                fail("idol_ids", "艺人 ID 不存在")
            artists.append({"type": "idols", "id": iid})
        if value.get("idol_id"):
            if not db.get(Idol, value["idol_id"]):
                fail("idol_id", "艺人 ID 不存在")
            if artists and not any(
                a["type"] == "idols" and a["id"] == value["idol_id"] for a in artists
            ):
                fail("idol_id", "idol_id 与所填艺人名称不一致")
        value["artists"] = list({(a["type"], a["id"]): a for a in artists}.values())
        value["idol_id"] = value.get("idol_id") or next(
            (a["id"] for a in value["artists"] if a["type"] == "idols"), None
        )
    if idol_name and kind not in {"groups", "idols"}:
        query = select(Idol).where(Idol.name == idol_name)
        matches = list({i.id: i for i in db.scalars(query).all()}.values())
        if len(matches) > 1 and value.get("group_id"):
            linked = db.scalars(
                select(Idol)
                .join(GroupMembership)
                .where(Idol.name == idol_name, GroupMembership.group_id == value["group_id"])
            ).all()
            if linked:
                matches = list({i.id: i for i in linked}.values())
        if len(matches) != 1:
            fail("idol_name", "艺人名称不存在或不唯一，请填写 idol_id")
        if value.get("idol_id") and value["idol_id"] != matches[0].id:
            fail("idol_name", "艺人名称与 idol_id 不一致")
        value["idol_id"] = matches[0].id
    value["source_info"] = {
        key: value.pop(key)
        for key in ["source_type", "source_url", "copyright_status", "copyright_note"]
        if value.get(key)
    }
    for field in ("release_date", "end_date"):
        if value.get(field):
            try:
                date.fromisoformat(value[field])
            except ValueError:
                fail(field, "日期请使用 YYYY-MM-DD 格式")
    if kind == "events":
        if (
            value.get("end_date")
            and value.get("release_date")
            and value["end_date"] < value["release_date"]
        ):
            fail("end_date", "结束日期不能早于开始日期")
        for field, related_kind in [("parent_event_id", "events"), ("album_id", "albums")]:
            if value.get(field) and not db.get(service.MODELS[related_kind], value[field]):
                fail(field, "关联对象不存在，请修正 ID")
    if value.get("image_url"):
        url = urlsplit(value["image_url"])
        if url.scheme not in {"http", "https"} or not url.netloc:
            fail("image_url", "封面图片须为有效的 HTTP(S) 地址")
    payload = EntityDraft.model_validate({k: v for k, v in value.items() if v != ""})
    if kind == "events":
        service.validate_event_location(payload)
    if any(not a.get("type") or not a.get("id") for a in payload.artists):
        fail("artists", "参与艺人须提供 type 与 id")
    if kind == "albums":
        if payload.idol_id and not db.get(Idol, payload.idol_id):
            fail("idol_id", "艺人 ID 不存在")
        if not payload.artists and payload.group_id:
            payload.artists = [{"type": "groups", "id": payload.group_id}]
        if payload.idol_id and not any(a.get("id") == payload.idol_id for a in payload.artists):
            payload.artists.append({"type": "idols", "id": payload.idol_id})
    return payload


def validate_rows(db: Session, kind: str, records: list[dict]) -> tuple[list, list]:
    if kind == "companies":
        from .company_import import validate_rows as validate_companies

        return validate_companies(db, records)
    if kind == "groups_idols":
        return validate_combined(db, records)
    valid, errors, seen, album_versions = [], [], set(), {}
    for index, record in enumerate(records, start=2):
        try:
            if None in record:
                fail("row", "数据列数超出表头；含逗号的内容请使用双引号包裹")
            value = {
                k: (v.date().isoformat() if isinstance(v, datetime) else v.isoformat())
                if isinstance(v, date)
                else str(v).strip()
                if v is not None
                else ""
                for k, v in record.items()
            }
            value["tags"] = [t.strip() for t in value.get("tags", "").split("|") if t.strip()]
            for field in ["versions", "aliases"]:
                if value.get(field):
                    value[field] = [t.strip() for t in value[field].split("|") if t.strip()]
            if kind == "albums" and value.get("version"):
                alias = [n.strip() for n in value.pop("version").split("|") if n.strip()]
                if value.get("versions") and value["versions"] != alias:
                    fail("versions", "version 与 versions 内容不一致，请只填写一个版本字段")
                value["versions"] = alias
            elif kind == "albums":
                value.pop("version", None)
            if value.get("artists"):
                import json

                try:
                    value["artists"] = json.loads(value["artists"])
                except ValueError:
                    fail("artists", "艺人归属必须为 JSON 数组")
            if kind == "cards":
                for side in ["front", "back"]:
                    value[f"{side}_source"] = {
                        "source_type": value.pop(f"{side}_source_type", "") or "ADMIN_UPLOAD",
                        "source_url": value.pop(f"{side}_source_url", ""),
                        "copyright_status": value.pop(f"{side}_copyright_status", "") or "UNKNOWN",
                    }
                payload = CardDraft.model_validate({k: v for k, v in value.items() if v != ""})
                row_errors = service.validate_relations(db, payload, True)
                from app.modules.catalog.tags import validate_names

                try:
                    validate_names(db, "CARD", payload.tags)
                except DomainError as exc:
                    row_errors.append({"field": "tags", "message": exc.message})
            else:
                payload = entity_draft(db, kind, value)
                row_errors = relation_errors(db, kind, payload)
                ownership = (
                    payload.group_id
                    if kind == "idols"
                    else artist_key(payload)
                    if kind != "groups"
                    else ""
                )
                identity = (ownership, payload.name)
                if kind == "albums" and payload.versions:
                    from .album_versions import version_key

                    versions = {version_key(n) for n in payload.versions}
                    if identity in seen and identity not in album_versions:
                        row_errors.append(
                            {"field": "name", "message": "同一专辑不能混合无版本行与版本行"}
                        )
                    elif versions & album_versions.get(identity, set()):
                        row_errors.append(
                            {"field": "versions", "message": "文件内存在同专辑重复版本"}
                        )
                    if not row_errors:
                        album_versions.setdefault(identity, set()).update(versions)
                elif identity in seen:
                    row_errors.append(
                        {"field": "name", "message": "文件内存在同组合同名记录，请移除重复行"}
                    )
                if not row_errors:
                    seen.add(identity)
            if not row_errors:
                valid.append({"row": index, "data": payload.model_dump()})
        except ValidationError as exc:
            row_errors = [
                {"field": ".".join(map(str, e["loc"])), "message": e["msg"]} for e in exc.errors()
            ]
        except DomainError as exc:
            row_errors = getattr(exc, "details", None) or [{"field": "row", "message": exc.message}]
        errors.extend(
            {"row": index, **e, "suggestion": e["message"], "can_skip": True} for e in row_errors
        )
    return valid, errors


def _combined_source_field(kind: str, field: str, data: dict) -> str:
    mapping = COMBINED_GROUP_FIELDS if kind == "groups" else COMBINED_IDOL_FIELDS
    if field.startswith("source_info."):
        field = field.removeprefix("source_info.")
    if kind == "idols" and field == "name":
        return "stage_name" if data.get("stage_name") else "idol_name"
    return next((source for source, target in mapping.items() if target == field), field)


def _combined_entity_values(kind: str, data: dict) -> dict:
    mapping = COMBINED_GROUP_FIELDS if kind == "groups" else COMBINED_IDOL_FIELDS
    values = {target: data.get(source, "") for source, target in mapping.items()}
    if kind == "idols":
        values["name"] = data.get("stage_name") or data.get("idol_name", "")
    for field in ("aliases", "tags"):
        values[field] = [
            item.strip() for item in str(values.get(field, "")).split("|") if item.strip()
        ]
    return values


def _combined_entity_draft(db: Session, kind: str, data: dict) -> tuple[EntityDraft | None, list]:
    try:
        return entity_draft(db, kind, _combined_entity_values(kind, data)), []
    except ValidationError as exc:
        return None, [
            {
                "field": _combined_source_field(kind, ".".join(map(str, error["loc"])), data),
                "message": error["msg"],
            }
            for error in exc.errors()
        ]
    except DomainError as exc:
        details = getattr(exc, "details", None) or [{"field": "row", "message": exc.message}]
        return None, [
            {
                **detail,
                "field": _combined_source_field(kind, detail.get("field", "row"), data),
            }
            for detail in details
        ]


def _combined_target(db: Session, kind: str, item_id: str):
    if not item_id:
        return None
    model = Group if kind == "groups" else Idol
    field = "group_id" if kind == "groups" else "idol_id"
    item = db.get(model, item_id)
    if not item:
        fail(field, "组合 ID 不存在" if kind == "groups" else "艺人 ID 不存在")
    status = service.entity_payload(db, kind, item)["status"]
    if status in {"PUBLISHED", "PENDING", "ARCHIVED"}:
        fail(field, f"当前状态为 {status}，请先下架、驳回或恢复草稿后再批量更新")
    return item


def _combined_targets(db: Session, data: dict) -> tuple[Group | None, Idol | None]:
    group = _combined_target(db, "groups", data.get("group_id", ""))
    idol = _combined_target(db, "idols", data.get("idol_id", ""))
    group_name = data.get("group_name", "")
    same_name_group = db.scalar(select(Group).where(Group.name == group_name))
    if group:
        if same_name_group and same_name_group.id != group.id:
            fail("group_name", "组合名称已被其他组合使用")
    elif same_name_group:
        fail("group_id", "组合已存在，请填写系统 group_id 使用更新逻辑")

    if not idol and group:
        idol_name = data.get("stage_name") or data.get("idol_name")
        existing_member = db.scalar(
            select(Idol)
            .join(GroupMembership, GroupMembership.idol_id == Idol.id)
            .where(GroupMembership.group_id == group.id, Idol.name == idol_name)
        )
        if existing_member:
            fail("idol_id", "该组合下艺人已存在，请填写系统 idol_id 使用更新逻辑")
    return group, idol


def _combined_company_plan(db: Session, data: dict, cache: dict) -> list[dict]:
    problems = []
    group_name = normalize_company_name(data.get("group_company", ""))
    idol_source_name = normalize_company_name(data.get("idol_company", ""))
    idol_inherited = not idol_source_name and bool(group_name)
    for prefix, name, field in [
        ("group", group_name, "group_company"),
        (
            "idol",
            idol_source_name or group_name,
            "group_company" if idol_inherited else "idol_company",
        ),
    ]:
        data[f"{prefix}_company_effective_name"] = name
        data[f"{prefix}_company_id"] = ""
        data[f"{prefix}_company_mode"] = "none"
        if prefix == "idol":
            data["idol_company_inherited"] = idol_inherited
        if not name:
            continue
        key = company_name_key(name)
        try:
            if key not in cache:
                cache[key] = resolve_company_name(db, name)
            company, mode = cache[key]
        except DomainError as exc:
            problems.append({"field": field, "message": exc.message})
            continue
        data[f"{prefix}_company_effective_name"] = (
            company.data.get("name", name) if company else name
        )
        data[f"{prefix}_company_id"] = company.id if company else ""
        data[f"{prefix}_company_mode"] = "reuse" if mode == "existing" else "create"
    return problems


def validate_combined(db: Session, records: list[dict]):
    prepared = []
    for index, row in enumerate(records, start=2):
        data = {
            key: (value.date().isoformat() if isinstance(value, datetime) else value.isoformat())
            if isinstance(value, date)
            else str(value or "").strip()
            for key, value in row.items()
        }
        problems = []
        if None in data:
            problems.append(
                {"field": "row", "message": "数据列数超出表头；含逗号的内容请使用双引号包裹"}
            )
        for field in ["group_code", "group_name", "idol_name"]:
            if not data.get(field):
                problems.append({"field": field, "message": "必填字段为空"})
        for field in ["group_id", "idol_id"]:
            if len(data.get(field, "")) > 64:
                problems.append({"field": field, "message": "长度不能超过 64 个字符"})
        if data.get("group_code") and data.get("group_name"):
            _, draft_errors = _combined_entity_draft(db, "groups", data)
            problems.extend(draft_errors)
        if data.get("idol_name"):
            _, draft_errors = _combined_entity_draft(db, "idols", data)
            problems.extend(draft_errors)
        prepared.append({"row": index, "data": data, "problems": problems})

    codes, group_names, group_ids, profiles, seen = {}, {}, {}, {}, set()
    for item in prepared:
        if item["problems"]:
            continue
        data = item["data"]
        name = data.get("stage_name") or data.get("idol_name", "")
        identity = (
            data.get("group_code", "").casefold(),
            data.get("idol_id", "").casefold() or name.casefold(),
        )
        if identity in seen:
            item["problems"].append({"field": "idol_name", "message": "文件内存在同组合重复艺人"})
        code = data.get("group_code", "").casefold()
        group_name = data.get("group_name", "").casefold()
        if code in codes and codes[code] != group_name:
            item["problems"].append(
                {"field": "group_name", "message": "同一 group_code 对应多个组合名称"}
            )
        if group_name in group_names and group_names[group_name] != code:
            item["problems"].append(
                {"field": "group_code", "message": "同一组合名称对应多个 group_code"}
            )
        group_id = data.get("group_id", "")
        if group_id and group_id in group_ids and group_ids[group_id] != code:
            item["problems"].append(
                {"field": "group_id", "message": "同一 group_id 对应多个 group_code"}
            )
        profile = profiles.setdefault(code, {})
        conflicts = []
        for field in COMBINED_GROUP_PROFILE_FIELDS:
            if not data.get(field) or not profile.get(field):
                continue
            if field == "group_company":
                matches = company_name_key(data[field]) == company_name_key(profile[field])
            else:
                matches = data[field] == profile[field]
            if not matches:
                conflicts.append(field)
        item["problems"].extend(
            {"field": field, "message": "同一组合的资料在文件内不一致"} for field in conflicts
        )
        if item["problems"]:
            continue
        for field in COMBINED_GROUP_PROFILE_FIELDS:
            if data.get(field):
                profile[field] = data[field]
        seen.add(identity)
        codes[code] = group_name
        group_names[group_name] = code
        if group_id:
            group_ids[group_id] = code

    valid, errors, company_cache = [], [], {}
    for item in prepared:
        index, data, problems = item["row"], item["data"], item["problems"]
        if not problems:
            code = data["group_code"].casefold()
            data.update(profiles[code])
            group_draft, draft_errors = _combined_entity_draft(db, "groups", data)
            problems.extend(draft_errors)
            idol_draft, draft_errors = _combined_entity_draft(db, "idols", data)
            problems.extend(draft_errors)
        if not problems:
            try:
                _combined_targets(db, data)
            except DomainError as exc:
                problems.extend(
                    getattr(exc, "details", None) or [{"field": "row", "message": exc.message}]
                )
        if not problems:
            problems.extend(_combined_company_plan(db, data, company_cache))
        if problems:
            errors += [
                {"row": index, **problem, "suggestion": problem["message"], "can_skip": True}
                for problem in problems
            ]
        else:
            data["group_operation"] = "update" if data.get("group_id") else "create"
            data["idol_operation"] = "update" if data.get("idol_id") else "create"
            data["group_data"] = group_draft.model_dump()
            data["idol_data"] = idol_draft.model_dump()
            valid.append({"row": index, "data": data})
    return valid, errors


def _combined_update_draft(
    db: Session, kind: str, item: Group | Idol, data: dict
) -> tuple[EntityDraft, bool]:
    current = {**service.entity_payload(db, kind, item), "company_ids": None}
    current.pop("official_accounts", None)
    before = EntityDraft.model_validate(current)
    mapping = COMBINED_GROUP_FIELDS if kind == "groups" else COMBINED_IDOL_FIELDS
    source_patch = {}
    for source, target in mapping.items():
        value = data.get(source, "")
        if value == "":
            continue
        if target in {"source_type", "source_url", "copyright_status", "copyright_note"}:
            source_patch[target] = value
        elif target in {"aliases", "tags"}:
            current[target] = [part.strip() for part in str(value).split("|") if part.strip()]
        else:
            current[target] = value
    if kind == "idols":
        current["name"] = data.get("stage_name") or data["idol_name"]
        current["real_name"] = data["idol_name"]
    if source_patch:
        current["source_info"] = {**(current.get("source_info") or {}), **source_patch}
    draft = EntityDraft.model_validate(current)
    return draft, draft.model_dump() != before.model_dump()


def import_combined(
    db: Session,
    actor,
    values: dict,
    new_groups: dict[str, Group] | None = None,
    new_companies: dict | None = None,
) -> dict:
    """Create or update both entities by ID, then guarantee their membership relation."""
    new_groups = new_groups if new_groups is not None else {}
    new_companies = new_companies if new_companies is not None else {}
    group_company = group_company_operation = None
    idol_company = idol_company_operation = None
    if values.get("group_company_effective_name"):
        group_company, group_company_operation = get_or_create_company(
            db,
            actor,
            values["group_company_effective_name"],
            new_companies,
        )
    if values.get("idol_company_effective_name"):
        idol_company, idol_company_operation = get_or_create_company(
            db,
            actor,
            values["idol_company_effective_name"],
            new_companies,
        )
    group_id = values.get("group_id", "")
    group_key = values["group_code"].casefold()
    if group_id:
        group = _combined_target(db, "groups", group_id)
        group_draft, changed = _combined_update_draft(db, "groups", group, values)
        if changed:
            service.save_entity(db, actor, "groups", group_draft, group.id)
        group_operation = "updated"
    else:
        group = new_groups.get(group_key)
        if not group:
            group_draft = EntityDraft.model_validate(values["group_data"])
            result = service.save_entity(db, actor, "groups", group_draft)
            group = db.get(Group, result["id"])
            new_groups[group_key] = group
        group_operation = "created"

    group_company_relation_operation = None
    if group_company:
        _, group_company_relation_operation = ensure_company_relation(
            db, actor, "groups", group.id, group_company.id
        )

    idol_id = values.get("idol_id", "")
    if idol_id:
        idol = _combined_target(db, "idols", idol_id)
        idol_draft, changed = _combined_update_draft(db, "idols", idol, values)
        if changed:
            service.save_entity(db, actor, "idols", idol_draft, idol.id)
        idol_operation = "updated"
    else:
        idol_draft = EntityDraft.model_validate(values["idol_data"])
        idol_draft.group_id = group.id
        result = service.save_entity(db, actor, "idols", idol_draft)
        idol = db.get(Idol, result["id"])
        idol_operation = "created"

    idol_company_relation_operation = None
    if idol_company:
        _, idol_company_relation_operation = ensure_company_relation(
            db, actor, "idols", idol.id, idol_company.id
        )

    membership = db.scalar(
        select(GroupMembership).where(
            GroupMembership.group_id == group.id,
            GroupMembership.idol_id == idol.id,
        )
    )
    membership_operation = "existing"
    if not membership:
        membership = GroupMembership(group_id=group.id, idol_id=idol.id)
        db.add(membership)
        db.flush()
        membership_operation = "created"
        service.audit(
            db,
            actor,
            "catalog.membership.import",
            "memberships",
            membership.id,
            None,
            service.row_data(membership),
        )
    if group_operation == "updated" and idol_operation == "updated":
        mode = "updated_group_and_idol"
    elif group_operation == "updated":
        mode = "updated_group_created_idol"
    elif idol_operation == "updated":
        mode = "created_group_updated_idol"
    else:
        mode = "created_group_and_idol"
    return {
        "created_id": idol.id,
        "group_id": group.id,
        "idol_id": idol.id,
        "membership_id": membership.id,
        "import_mode": mode,
        "group_operation": group_operation,
        "idol_operation": idol_operation,
        "membership_operation": membership_operation,
        "group_company_id": group_company.id if group_company else "",
        "group_company_operation": group_company_operation or "none",
        "group_company_relation_operation": group_company_relation_operation or "none",
        "idol_company_id": idol_company.id if idol_company else "",
        "idol_company_operation": idol_company_operation or "none",
        "idol_company_relation_operation": idol_company_relation_operation or "none",
        "idol_company_inherited": bool(values.get("idol_company_inherited")),
    }
